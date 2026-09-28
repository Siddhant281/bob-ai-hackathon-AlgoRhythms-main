# docs.py
# Placeholder for document generation utilities (python-docx).

from docx import Document
from docx.shared import Pt
import io


def generate_sho_brief(data: dict) -> bytes:
    """Generate a SHO Brief Word document and return its bytes."""
    doc = Document()
    doc.add_heading("SHO Brief", level=1)

    for key, value in data.items():
        p = doc.add_paragraph()
        run = p.add_run(f"{key}: ")
        run.bold = True
        run.font.size = Pt(11)
        p.add_run(str(value))

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
